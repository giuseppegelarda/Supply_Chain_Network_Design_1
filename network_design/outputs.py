"""Excel and chart helpers, the sensitivity report and tornado chart, the CO2 report and the charts of
Output/."""
import os
from io import BytesIO
import pandas as pd
import numpy as np
from network_design.core import (_demand_dict, _config_display, _print_header, ROW_TOTAL_CO2_COST, ROW_UNITS_SOLD,
                                 ROW_LOST_SALES, ROW_FILL_RATE, ROW_CO2_EMISSIONS, CO2_EMISSION_SOURCES,
                                 LOST_SALES_TOLERANCE, FIXED_COST_KEYS, NO_WAREHOUSE)
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter


def _save_figure(path):
    """Saves the current figure to path (creating its folder if needed), closes it and logs the path."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    print(f"Saved: {path}")


def _figure_png(fig):
    """The figure as PNG bytes for an Excel sheet (closes it)."""
    image = BytesIO()
    fig.savefig(image, format="png", dpi=100, bbox_inches="tight")
    plt.close(fig)
    image.seek(0)
    return image


def _format_value(value):
    """Readable value: 2,546,875 / 146.1 / 0.1458 (4 significant digits below 1,000)."""
    if abs(value) >= 1000:
        return f"{value:,.0f}"
    return f"{value:.4g}"


SENSITIVITY_INVENTORY_NOTE = (
    "Note: the inventory cost is not in the solver's objective, so when a cost falls the solver can add a third "
    "Plant -> Warehouse route and the reported expected profit can fall slightly (up to about 32,000 EUR/year, "
    "0.13%, found with Transport Shanghai to Mumbai at -20% in a run with all the values for every parameter); it "
    "never changes the recommended network.")
# White background behind chart labels, so that no line crosses them
LABEL_BOX = {"facecolor": "white", "edgecolor": "none", "alpha": 0.85, "boxstyle": "round,pad=0.15"}


def _network_colors(labels):
    """Colors of the networks in the charts, the same in every analysis (tab10 in network order; no warehouse:
    gray)."""
    palette = plt.get_cmap("tab10").colors
    named = [label for label in labels if label != NO_WAREHOUSE]
    colors = {label: palette[i % len(palette)] for i, label in enumerate(named)}
    colors[NO_WAREHOUSE] = "gray"
    return colors


# Number formats shared by the Excel reports: (Excel format, Python format of the same text, for the column widths)
FORMAT_EUR = ('#,##0 "EUR"', "{:,.0f} EUR")
FORMAT_EUR_PER_UNIT = ('0.0000 "EUR/unit"', "{:.4f} EUR/unit")
FORMAT_KG = ('#,##0 "kg"', "{:,.0f} kg")
FORMAT_KG_PER_UNIT = ('0.0000 "kg/unit"', "{:.4f} kg/unit")
FORMAT_UNITS = ('#,##0 "units"', "{:,.0f} units")
FORMAT_INTEGER = ("0", "{:.0f}")
FORMAT_PERCENT = ("0.00%", "{:.2%}")
FORMAT_PERCENT_1 = ("0.0%", "{:.1%}")


def _format_table(ws, df, start_row, start_col, cell_format):
    """Number format of each number of a written table, cell_format(first-column value, column) -> (Excel format,
    Python format) or None; columns sized to the text."""
    widths = [len(str(column)) for column in df.columns]
    for r, row in enumerate(df.itertuples(index=False)):
        for c, (column, value) in enumerate(zip(df.columns, row)):
            formats = cell_format(row[0], column) if isinstance(value, (int, float, np.number)) else None
            if formats:
                ws.cell(row=start_row + r + 2, column=start_col + c + 1).number_format = formats[0]
            widths[c] = max(widths[c], len(formats[1].format(value) if formats else str(value)))
    for c, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + c + 1)].width = width + 2


def _write_excel(output_dir, filename, sheets, message, notes=None, cell_format=None, images=None):
    """Writes {sheet: DataFrame or list of DataFrames side by side, one empty column apart} without the index
    (notes: {sheet: text} in A1, the tables from row 3; cell_format: see _format_table; images: {sheet: PNG}
    right of the tables), logs the path."""
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, tables in sheets.items():
            start_col = 0
            for df in (tables if isinstance(tables, list) else [tables]):
                df.to_excel(writer, sheet_name=sheet, index=False, startrow=2 if notes else 0, startcol=start_col)
                if cell_format:
                    _format_table(writer.sheets[sheet], df, 2 if notes else 0, start_col, cell_format)
                start_col += len(df.columns) + 1
            if notes:
                writer.sheets[sheet].cell(row=1, column=1, value=notes[sheet])
            if images and sheet in images:
                writer.sheets[sheet].add_image(ExcelImage(images[sheet]), f"{get_column_letter(start_col + 1)}1")
    print(f"{message}: {path}")
    return path


def _change_text(change, base_value):
    """A recommended network change as text, e.g. "above 4,141,929 (+107.1%): Kolkata"."""
    value, below, above = change
    side, network = ("below", below) if value < base_value else ("above", above)
    return f"{side} {_format_value(value)} ({(value - base_value) / abs(base_value):+.1%}): {network}"


def plot_sensitivity_tornado(df_summary, low_label, high_label, base_expected, n_scenarios,
                             filename="sensitivity_tornado.png", min_co2_change_pct=0.1):
    """Tornado chart of the expected profit change at the ends of each parameter's range."""
    df_plot = df_summary.iloc[::-1]  # the largest impact at the top
    recommended = base_expected.index[0]
    fig, ax = plt.subplots(figsize=(11, max(4, 0.6 * len(df_plot))))
    y_pos = np.arange(len(df_plot))

    def change(r, label):
        network = r[f"Recommended Network at {label}"]
        if network != recommended:
            return f"-> {network}"
        return "mix" if abs(r[f"Expected CO2 Change at {label} (%)"]) > min_co2_change_pct else "-"

    for label, color in ((low_label, "tab:red"), (high_label, "tab:blue")):
        values = df_plot[f"Expected Profit Change at {label} (%)"]
        bars = ax.barh(y_pos, values, color=color)
        for bar, (_, r), value in zip(bars, df_plot.iterrows(), values):
            if change(r, label) == "mix":
                bar.set_hatch("//")
                bar.set_edgecolor("white")
                bar.set_linewidth(0)  # a white border would cover the thin bars
            if value:  # value at the outer end of the bar
                ax.annotate(f"{value:+.{2 if abs(value) < 10 else 1}f}%", (value, bar.get_y() + bar.get_height() / 2),
                            xytext=(3 if value > 0 else -3, 0), textcoords="offset points",
                            ha="left" if value > 0 else "right", va="center", fontsize=8)
    ax.margins(x=0.15)  # room for the values of the longest bars
    for y, (_, r) in zip(y_pos, df_plot.iterrows()):
        ax.text(1.02, y, f"{change(r, low_label)}  |  {change(r, high_label)}", transform=ax.get_yaxis_transform(),
                va="center", fontsize=9)
    ax.text(1.02, 1.0, f"Change vs base\n({low_label}  |  {high_label})", transform=ax.transAxes, va="bottom",
            fontsize=9, fontweight="bold")
    # Legend handles built here: a bar container's own handle takes the style of its first bar, which can be hatched
    handles = [plt.Rectangle((0, 0), 1, 1, color=color, label=f"Parameter {label}")
               for label, color in ((low_label, "tab:red"), (high_label, "tab:blue"))]
    handles.append(plt.Rectangle((0, 0), 1, 1, facecolor="white", edgecolor="black", hatch="//",
                                 label="Production mix changes"))
    ax.axvline(0, color="black", linewidth=0.8, zorder=0.5)  # behind the bars: on top it would hide the thinnest
    ax.set_yticks(y_pos)
    ax.set_yticklabels(df_plot["Parameter"])
    ax.set_xlabel("Expected Profit % Change", fontsize=12)
    ax.set_title(f"Sensitivity of the Recommended Network (Tornado Chart): Impact on Expected Profit\n(each parameter "
                 f"changed by ±{high_label.lstrip('+')} on its own, networks kept fixed on {n_scenarios} demand "
                 f"scenarios)", fontsize=13, pad=30)  # above the column header, which it would cover
    ax.legend(handles=handles, loc="lower right")
    ax.grid(axis="x", alpha=0.3)
    _save_figure(filename)


def _sensitivity_chart(result, colors, n_scenarios):
    """Chart of one parameter (expected profit of every network, CO2 of the recommended) as PNG bytes."""
    spec, curve, base_value = result["spec"], result["curve"], result["base_value"]
    x, recommended = curve["Parameter Value"], curve["Recommended Network"]
    networks = [label for label in colors if label != NO_WAREHOUSE or (recommended == label).any()]
    with_co2 = spec["key"] not in FIXED_COST_KEYS  # CAPEX, OPEX, discount rate: flows and CO2 never change
    if with_co2:
        fig, (ax, ax_co2) = plt.subplots(2, 1, figsize=(10, 8), sharex=True, gridspec_kw={"height_ratios": [3, 2]})
        axes = (ax, ax_co2)
    else:
        fig, ax = plt.subplots(figsize=(10, 4.8))
        axes = (ax,)
    for label in networks:
        ax.plot(x, curve[f"Expected Profit - {label}"] / 1e6, marker="o", markersize=3, linewidth=1.5,
                color=colors[label], label=label)
    for label in recommended.unique():
        chosen = recommended == label
        ax.scatter(x[chosen], curve.loc[chosen, f"Expected Profit - {label}"] / 1e6, s=90, color=colors[label],
                   edgecolors="black", linewidths=1.2, zorder=3)
        if with_co2:
            ax_co2.scatter(x[chosen], curve.loc[chosen, "Expected CO2 (kg/year)"], s=60, color=colors[label],
                           edgecolors="black", zorder=3, label=label)
    within = sorted(c for c in result["changes"] if x.iloc[0] <= c[0] <= x.iloc[-1])
    for j, (value, below, above) in enumerate(within):  # labels at alternating heights, as they can be close
        for axis in axes:
            axis.axvline(value, color="red", linestyle="--", alpha=0.6)
        ax.annotate(f"{_format_value(value)}: {below} -> {above}", (value, 0.97 - 0.07 * (j % 2)),
                    xycoords=ax.get_xaxis_transform(), xytext=(4, 0), textcoords="offset points", ha="left",
                    va="top", fontsize=9, fontweight="bold", color="red", bbox=LABEL_BOX)
    for axis in axes:
        axis.axvline(base_value, color="black", linestyle=":", linewidth=1.5)
    ax.text(base_value, 1.01, f"base {_format_value(base_value)}", transform=ax.get_xaxis_transform(), ha="center",
            va="bottom", fontsize=9)
    beyond = [c for c in result["changes"] if not x.iloc[0] <= c[0] <= x.iloc[-1]]
    ax.set_title(f"Sensitivity of the Recommended Network: {spec['label']}\n(every network kept fixed on "
                 f"{n_scenarios} demand scenarios; large markers: recommended network"
                 + ("; red dashed lines: exact changes" if within else "") + ")"
                 + (f"\nbeyond the tested range: {', '.join(_change_text(c, base_value) for c in beyond)}"
                    if beyond else ""), fontsize=12, pad=20)
    handles, _ = ax.get_legend_handles_labels()
    handles.append(plt.Line2D([], [], marker="o", linestyle="", markersize=9, markerfacecolor="white",
                              markeredgecolor="black", label="recommended network"))
    ax.legend(handles=handles, title="Network (fixed)", loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=9)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:,.2f}"))
    ax.set_ylabel("Expected profit, network fixed (EUR M/year)", fontsize=11)
    ax.grid(alpha=0.3)
    if with_co2:
        ax_co2.legend(title="Recommended network", loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=9)
        ax_co2.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:,.0f}"))
        ax_co2.set_ylabel("Expected CO2 of the\nrecommended network (kg/year)", fontsize=11)
        ax_co2.grid(alpha=0.3)
    if x.abs().max() >= 1000:  # e.g. CAPEX: 2,000,000 instead of 2.0 with "1e6"
        axes[-1].xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:,.0f}"))
    axes[-1].set_xlabel(f"{spec['label']} ({spec['unit']})", fontsize=11)
    return _figure_png(fig)


# Number format of a parameter's value in sensitivity_results.xlsx, by its unit: (Excel format, Python format)
SENSITIVITY_UNIT_FORMATS = {"EUR/unit": FORMAT_EUR_PER_UNIT, "EUR": FORMAT_EUR,
                            "EUR/year": ('#,##0 "EUR/year"', "{:,.0f} EUR/year"),
                            "EUR/ton CO2": ('#,##0 "EUR/ton CO2"', "{:,.0f} EUR/ton CO2"),
                            "per year": FORMAT_PERCENT}


def _sensitivity_format(column, unit):
    """(Excel format, Python format) of a number of sensitivity_results.xlsx, by its column and parameter unit."""
    if column.endswith("(%)"):  # already in percentage points
        return '0.00"%"', "{:.2f}%"
    if column.startswith(("Expected Profit - ", "Advantage")):
        return FORMAT_EUR
    if column == "Expected CO2 (kg/year)":
        return FORMAT_KG
    return SENSITIVITY_UNIT_FORMATS[unit]  # a value of the parameter


def export_sensitivity_results(results, df_summary, output_dir, n_scenarios):
    """Writes sensitivity_results.xlsx: Summary and one sheet with its chart per detailed case."""
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "sensitivity_results.xlsx")
    labels = [column.removeprefix("Expected Profit - ") for column in results[0]["curve"].columns
              if column.startswith("Expected Profit - ")]
    colors = _network_colors(labels)
    units = {result["spec"]["label"]: result["spec"]["unit"] for result in results}
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_summary.to_excel(writer, sheet_name="Summary", index=False)
        _format_table(writer.sheets["Summary"], df_summary, 0, 0,
                      lambda label, column: _sensitivity_format(column, units[label]))
        writer.sheets["Summary"].cell(row=len(df_summary) + 3, column=1, value=SENSITIVITY_INVENTORY_NOTE)
        for k, result in enumerate([result for result in results if result["detail"]], start=1):
            sheet_name, spec = f"S{k:02d}", result["spec"]
            curve = result["curve"].rename(columns={"Parameter Value": f"{spec['label']} ({spec['unit']})"})
            curve.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_table(writer.sheets[sheet_name], curve, 0, 0,
                          lambda label, column, unit=spec["unit"]: _sensitivity_format(column, unit))
            chart = ExcelImage(_sensitivity_chart(result, colors, n_scenarios))
            writer.sheets[sheet_name].add_image(chart, f"{get_column_letter(curve.shape[1] + 2)}1")
    print(f"Sensitivity analysis report generated: {output_path}")
    return output_path


def _analysis_format(label, column):
    """(Excel format, Python format) of a number of the N_YEARS, capacity and CV reports, by its column."""
    if column == "Value":  # Settings sheet, by row
        return FORMAT_PERCENT if label == "Discount Rate" else FORMAT_INTEGER
    if column.startswith("Share Optimal"):
        return FORMAT_PERCENT_1
    if "Break-even" in column:
        return '0.00 "years"', "{:.2f} years"
    if column.startswith(("Capacity", "Mean Total Demand")):
        return FORMAT_UNITS
    if column in ("CV", "CAPEX Annualization Factor"):
        return ("0.00", "{:.2f}") if column == "CV" else ("0.0000", "{:.4f}")
    if column == "N_YEARS" or column.startswith(("Feasible", "N. ", "Most Frequently")):
        return FORMAT_INTEGER
    return FORMAT_EUR  # warehouse costs, profits, advantages, gaps, WS, RP, EEV, EVPI, VSS


def report_co2_emissions(fixed_evaluations, n_feasible, output_dir="Output"):
    """Average kg of CO2 of each fixed configuration on the common scenarios: Excel report and chart."""
    feasible_sets = [set(ev["df_fixed"]["scenario"]) if not ev["df_fixed"].empty else set()
                     for ev in fixed_evaluations.values()]
    common = sorted(set.intersection(*feasible_sets)) if feasible_sets else []
    if not common:
        print("Skipped: CO2 emissions comparison (no scenario evaluated for every configuration)")
        return

    summary_rows, per_scenario = [], {}
    for ev in fixed_evaluations.values():
        label = _config_display(ev["open_warehouses"])
        df = ev["df_fixed"].set_index("scenario").loc[common]
        summary_rows.append({
            "Configuration": label,
            "N. Scenarios Used (common to all configurations)": len(common),
            "N. Scenarios With Lost Sales": int((df[ROW_LOST_SALES] > LOST_SALES_TOLERANCE).sum()),
            "Average Fill Rate": df[ROW_FILL_RATE].mean(),
            "Total CO2 Emissions (kg)": df[ROW_CO2_EMISSIONS].mean(),
            **{f"{source} (kg)": df[row].mean() for source, row in CO2_EMISSION_SOURCES.items()},
            "CO2 Emissions per Unit Sold (kg/unit)": (df[ROW_CO2_EMISSIONS] / df[ROW_UNITS_SOLD]).mean(),
            "Total CO2 Offsetting Cost (EUR)": df[ROW_TOTAL_CO2_COST].mean(),
        })
        per_scenario[f"{label} (kg)"] = df[ROW_CO2_EMISSIONS]
    df_summary = pd.DataFrame(summary_rows)

    all_scenarios = sorted(set.union(*feasible_sets, *(set(ev["discarded"]) for ev in fixed_evaluations.values())))
    lost_sales_sets = {_config_display(ev["open_warehouses"]):
                       set(ev["df_fixed"].loc[ev["df_fixed"][ROW_LOST_SALES] > LOST_SALES_TOLERANCE, "scenario"])
                       if not ev["df_fixed"].empty else set() for ev in fixed_evaluations.values()}
    df_scen = pd.DataFrame({
        "Scenario": all_scenarios,
        "Used": ["Yes" if s in common else "No" for s in all_scenarios],
        "Networks With Lost Sales": [", ".join(label for label, lost in lost_sales_sets.items() if s in lost)
                                     or "-" for s in all_scenarios],
    })

    def cell_format(label, column):  # by the unit in the column name
        if column.endswith("(kg/unit)"):
            return FORMAT_KG_PER_UNIT
        if column.endswith("(kg)"):
            return FORMAT_KG
        if column.endswith("(EUR)"):
            return FORMAT_EUR
        return FORMAT_PERCENT if column == "Average Fill Rate" else FORMAT_INTEGER  # scenario numbers and counts

    _write_excel(output_dir, "co2_emissions.xlsx",
                 {"Summary": df_summary,
                  "Per_Scenario": pd.DataFrame(per_scenario).rename_axis("Scenario").reset_index(),
                  "Scenarios": df_scen},
                 "CO2 emissions report generated", cell_format=cell_format)

    fig, ax = plt.subplots(figsize=(max(7, 3 * len(df_summary)) + 2.5, 6.5))
    bottom = np.zeros(len(df_summary))
    for source, color in zip(CO2_EMISSION_SOURCES, ["tab:blue", "tab:orange", "tab:green", "tab:purple"]):
        values = df_summary[f"{source} (kg)"].to_numpy()
        ax.bar(df_summary["Configuration"], values, bottom=bottom, label=source, color=color, edgecolor="black",
               width=0.6)
        bottom += values
    for idx, row in df_summary.iterrows():
        ax.text(idx, row["Total CO2 Emissions (kg)"],
                f"{row['Total CO2 Emissions (kg)']:,.0f} kg/year\n"
                f"{row['CO2 Emissions per Unit Sold (kg/unit)']:.4f} kg/unit sold\n"
                f"fill rate {row['Average Fill Rate']:.1%}\n"
                f"lost sales in {row['N. Scenarios With Lost Sales']}/{len(common)} scenarios",
                ha="center", va="bottom", fontsize=9, fontweight="bold")
    ax.set_ylim(0, bottom.max() * 1.3)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:,.0f}"))
    ax.set_ylabel("Average CO2 emitted (kg/year)", fontsize=12)
    ax.set_xlabel("Configuration (fixed network)", fontsize=12)
    ax.set_title("Average CO2 Emissions by Configuration\n"
                 f"(average over the same {len(common)} of {n_feasible} scenarios for every configuration)",
                 fontsize=13)
    ax.legend(title="Emission source", loc="upper left", bbox_to_anchor=(1.01, 1))
    ax.grid(axis="y", linestyle="--", alpha=0.7)
    _save_figure(os.path.join(output_dir, "co2_emissions_by_configuration.png"))


def plot_production_by_plant(x_sol, plants, warehouses, recommended_info, plant_capacity, output_dir="Output"):
    """Units produced per plant by warehouse, in the recommended network's representative scenario, with the
    capacity of each plant."""
    fig, ax = plt.subplots(figsize=(8, 5))
    bottom = np.zeros(len(plants))
    colors = ["tab:blue", "tab:orange", "tab:green", "tab:red", "tab:purple"]
    for idx, k in enumerate(warehouses):
        values = np.array([x_sol.get((i, k), 0.0) or 0.0 for i in plants])
        ax.bar(plants, values, bottom=bottom, label=k, edgecolor="black", color=colors[idx % len(colors)])
        bottom += values
    for idx, total in enumerate(bottom):
        # Inside the bar, below its top: a production close to the capacity would put the label on the line
        ax.annotate(f"{total:,.0f}", (idx, total), xytext=(0, -4), textcoords="offset points", ha="center",
                    va="top", fontsize=10, fontweight="bold", bbox=LABEL_BOX)
    for idx, plant in enumerate(plants):
        ax.hlines(plant_capacity[plant], idx - 0.4, idx + 0.4, colors="black", linestyles="--", linewidth=1.5,
                  zorder=3)
        ax.annotate(f"Capacity {plant_capacity[plant] / 1e6:.1f} M", (idx - 0.4, plant_capacity[plant]),
                    xytext=(0, 3), textcoords="offset points", ha="left", va="bottom", fontsize=9, zorder=4,
                    bbox=LABEL_BOX)
    ax.set_ylim(0, max(max(plant_capacity.values()), bottom.max()) * 1.1)  # room for the capacity labels
    ax.set_xlabel("Plant", fontsize=12)
    ax.set_ylabel("Units produced and shipped", fontsize=12)
    ax.set_title("Production per Plant, by Destination Warehouse\n"
                 f"(Recommended configuration: {recommended_info['config_label']} | "
                 f"Representative scenario {recommended_info['scenario']}, network fixed)", fontsize=13)
    ax.legend(title="Warehouse")
    ax.grid(axis="y", linestyle="--", alpha=0.7)
    _save_figure(os.path.join(output_dir, "production_by_plant.png"))


def plot_demand_with_mean(df_demand, markets, base_demand, warehouse_capacity, plant_capacity, output_dir="Output"):
    """Demand scenarios of each market and of the total, with mean, base demand and capacities."""
    df_plot = df_demand.assign(scenario=range(1, len(df_demand) + 1))
    colors = ["tab:green", "tab:red", "black", "tab:blue", "tab:orange", "tab:purple"]
    millions = plt.FuncFormatter(lambda v, _: f"{v / 1e6:.1f}")

    def draw(ax, values, base_value, color, label):
        ax.plot(df_plot["scenario"], values, color=color, label=label)
        ax.axhline(values.mean(), color=color, linestyle="--", linewidth=1.5,
                   label=f"Mean of {len(values)} scenarios: {values.mean() / 1e6:.2f} M")
        ax.axhline(base_value, color="dimgray", linestyle=":", linewidth=2,
                   label=f"Base demand: {base_value / 1e6:.2f} M")
        ax.yaxis.set_major_formatter(millions)
        ax.grid(True)

    fig, axes = plt.subplots(len(markets) + 1, 1, figsize=(15, 3 * (len(markets) + 1)), sharex=True)
    for i, mkt in enumerate(markets):
        if i > 0:
            axes[i].sharey(axes[0])  # same scale for every market, so their spread can be compared
        draw(axes[i], df_plot[mkt], base_demand[mkt], colors[i % len(colors)], mkt)
        axes[i].set_ylabel(f"{mkt}\n(M units/year)", fontsize=11)

    ax_total = axes[-1]
    draw(ax_total, df_plot[markets].sum(axis=1), sum(base_demand[mkt] for mkt in markets), "tab:purple",
         "Total demand")
    by_capacity = {}  # warehouses grouped by capacity, one threshold line per capacity value
    for warehouse, capacity in warehouse_capacity.items():
        by_capacity.setdefault(capacity, []).append(warehouse)
    for capacity, names in by_capacity.items():
        ax_total.axhline(capacity, color="tab:red", linestyle="-.", linewidth=1.5,
                         label=f"Capacity of one warehouse ({', '.join(names)}): {capacity / 1e6:.1f} M")
    ax_total.axhline(plant_capacity, color="black", linestyle="-.", linewidth=1.5,
                     label=f"Total plant capacity: {plant_capacity / 1e6:.1f} M")
    ax_total.set_ylabel("Total demand\n(M units/year,\nown scale)", fontsize=11)
    for ax in axes:  # legends outside the panels, so they never cover the scenarios
        ax.legend(loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=9)
    ax_total.set_xlim(1, len(df_plot))
    ax_total.set_xlabel("Scenario", fontsize=12)
    fig.suptitle("Market Demand Scenarios, Mean and Base Demand\n(the total demand panel has its own scale; above "
                 "the capacity of one warehouse, a network with that warehouse alone loses sales)", y=1.01, fontsize=14)
    _save_figure(os.path.join(output_dir, "demand_with_mean.png"))


def generate_all_plots(df_scenarios, markets, x_recommended, recommended_info, plants, warehouses, params,
                       output_dir="Output"):
    """Generates all the chart reports (params: the input data, for the base demand and the capacities)."""
    _print_header("GENERATING GRAPHICAL REPORTS")
    plot_production_by_plant(x_recommended, plants, warehouses, recommended_info,
                             params["cap_pl"].iloc[:, 0].to_dict(), output_dir=output_dir)
    plot_demand_with_mean(df_scenarios, markets, _demand_dict(params["demand"]), params["cap_wh"].iloc[:, 0].to_dict(),
                          float(params["cap_pl"].iloc[:, 0].sum()), output_dir=output_dir)
    print("\nAll plots generated successfully!")
